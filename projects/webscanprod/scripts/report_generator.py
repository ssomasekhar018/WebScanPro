from pathlib import Path
import pandas as pd
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import inch
from docx import Document
from docx.shared import Inches


def webscanprod_root() -> Path:
    return Path(__file__).resolve().parent.parent


BASE_DIR = webscanprod_root()
DATA_CSV = BASE_DIR / "data" / "processed" / "findings.csv"
REPORTS_DIR = BASE_DIR / "reports"


def ensure_reports_dir():
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)


def generate_markdown(df: pd.DataFrame):
    ensure_reports_dir()
    md_path = REPORTS_DIR / "webscan_summary.md"
    total = len(df)
    severity_counts = df["severity"].value_counts().to_dict() if total else {}

    lines = []
    lines.append("# WebScanProd — Vulnerability Summary")
    lines.append("")
    lines.append(f"Total Findings: {total}")
    lines.append("")
    lines.append("## Severity Breakdown")
    for sev in ["Critical", "High", "Medium", "Low"]:
        lines.append(f"- {sev}: {severity_counts.get(sev, 0)}")
    lines.append("")
    lines.append("## Findings")
    if total:
        for _, r in df.iterrows():
            lines.append(f"### {r['vuln_id']} — {r['vuln_name']} ({r['severity']})")
            lines.append(f"- Affected: `{r['affected_url']}`")
            if r.get("impact"):
                lines.append(f"- Impact: {r['impact']}")
            lines.append("- Description:")
            lines.append(f"  {r['description']}")
            lines.append("- Mitigation:")
            lines.append(f"  {r['recommendation']}")
            lines.append("")
    else:
        lines.append("No findings available.")

    md_path.write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote Markdown summary: {md_path}")


def generate_pdf(df: pd.DataFrame):
    ensure_reports_dir()
    pdf_path = REPORTS_DIR / "vulnerability_report.pdf"
    c = canvas.Canvas(str(pdf_path), pagesize=A4)
    width, height = A4

    def write_line(text, x=inch, y_offset=[height - inch]):
        # y_offset as single-item list for mutability in closure
        c.drawString(x, y_offset[0], text)
        y_offset[0] -= 14

    write_line("WebScanProd — Vulnerability Report")
    write_line("")
    write_line(f"Total Findings: {len(df)}")
    sev_counts = df["severity"].value_counts().to_dict() if len(df) else {}
    write_line("Severity Breakdown:")
    for sev in ["Critical", "High", "Medium", "Low"]:
        write_line(f"- {sev}: {sev_counts.get(sev, 0)}")
    write_line("")
    write_line("Findings:")

    for _, r in df.iterrows():
        write_line(f"{r['vuln_id']} — {r['vuln_name']} ({r['severity']})")
        write_line(f"Affected: {r['affected_url']}")
        if r.get("impact"):
            write_line(f"Impact: {r['impact']}")
        write_line("Description:")
        write_line(r['description'][:200])
        write_line("Mitigation:")
        write_line(r['recommendation'][:200])
        write_line("")

    c.showPage()
    c.save()
    print(f"Wrote PDF report: {pdf_path}")


def generate_docx(df: pd.DataFrame):
    ensure_reports_dir()
    docx_path = REPORTS_DIR / "vulnerability_report.docx"
    doc = Document()
    doc.add_heading('WebScanProd — Vulnerability Report', 0)
    doc.add_paragraph(f"Total Findings: {len(df)}")
    sev_counts = df["severity"].value_counts().to_dict() if len(df) else {}
    doc.add_paragraph("Severity Breakdown:")
    for sev in ["Critical", "High", "Medium", "Low"]:
        doc.add_paragraph(f"- {sev}: {sev_counts.get(sev, 0)}")

    doc.add_heading('Findings', level=1)
    table = doc.add_table(rows=1, cols=7)
    hdr = table.rows[0].cells
    hdr[0].text = 'ID'
    hdr[1].text = 'Name'
    hdr[2].text = 'Severity'
    hdr[3].text = 'Affected URL'
    hdr[4].text = 'Impact'
    hdr[5].text = 'Description'
    hdr[6].text = 'Recommendation'

    for _, r in df.iterrows():
        row = table.add_row().cells
        row[0].text = str(r.get('vuln_id', ''))
        row[1].text = str(r.get('vuln_name', ''))
        row[2].text = str(r.get('severity', ''))
        row[3].text = str(r.get('affected_url', ''))
        row[4].text = str(r.get('impact', ''))
        row[5].text = str(r.get('description', ''))
        row[6].text = str(r.get('recommendation', ''))

    doc.save(str(docx_path))
    print(f"Wrote DOCX report: {docx_path}")


def main():
    if not DATA_CSV.exists():
        print(f"Processed dataset not found: {DATA_CSV}. Run scan_runner.py first.")
        return
    df = pd.read_csv(DATA_CSV)
    generate_markdown(df)
    generate_pdf(df)
    generate_docx(df)


if __name__ == "__main__":
    main()

